"""
Records Manager for Patient EHR.
Reads and writes structured patient records to/from Excel (DataSet/records.xlsx)
using openpyxl, maintaining a safe runtime copy in data_storage.
"""

import shutil
from pathlib import Path
from typing import Any, Dict, List, Optional
import openpyxl

from src.config import EXCEL_RECORDS_PATH, RUNTIME_EXCEL_PATH


class RecordsManager:
    """
    Manages structured patient records stored in Excel.
    Ensures safe read/write operations without modifying original training dataset unless desired.
    """

    EXPECTED_COLUMNS = ["Phone_number", "Email", "Name", "Age", "Gender", "Address", "Summary"]

    def __init__(self, source_path: Path = EXCEL_RECORDS_PATH, runtime_path: Path = RUNTIME_EXCEL_PATH):
        self.source_path = Path(source_path)
        self.runtime_path = Path(runtime_path)
        self._initialize_runtime_file()

    def _initialize_runtime_file(self, force_reset: bool = False) -> None:
        """Initializes the working Excel file by copying from source if not present or on force_reset."""
        self.runtime_path.parent.mkdir(parents=True, exist_ok=True)
        if not self.runtime_path.exists() or force_reset:
            if self.source_path.exists():
                shutil.copy2(self.source_path, self.runtime_path)
            else:
                # Create a blank sheet with headers
                wb = openpyxl.Workbook()
                ws = wb.active
                ws.title = "Sheet1"
                ws.append(self.EXPECTED_COLUMNS)
                wb.save(self.runtime_path)

    def reset_to_default(self) -> None:
        """Resets runtime records back to the original dataset state."""
        self._initialize_runtime_file(force_reset=True)

    def get_all_patients(self) -> List[Dict[str, Any]]:
        """Returns all patient records as a list of dictionaries."""
        wb = openpyxl.load_workbook(self.runtime_path)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            return []

        headers = [str(h).strip() if h is not None else f"col_{idx}" for idx, h in enumerate(rows[0])]
        patients = []

        for row in rows[1:]:
            if not any(row):
                continue
            patient_dict = {}
            for col_idx, h in enumerate(headers):
                val = row[col_idx] if col_idx < len(row) else None
                patient_dict[h] = str(val).strip() if val is not None else ""
            patients.append(patient_dict)

        return patients

    def get_patient(self, identifier: str) -> Optional[Dict[str, Any]]:
        """
        Finds a patient by Name or Phone number (case-insensitive fuzzy/exact match).
        """
        identifier = str(identifier).strip().lower()
        if not identifier:
            return None

        patients = self.get_all_patients()
        for p in patients:
            p_name = p.get("Name", "").lower()
            p_phone = p.get("Phone_number", "").lower().replace("-", "").replace(" ", "").replace("+", "")
            clean_id = identifier.replace("-", "").replace(" ", "").replace("+", "")

            # Exact or substring match
            if identifier == p_name or identifier in p_name or p_name in identifier:
                return p
            if clean_id and clean_id in p_phone:
                return p

        return None

    def add_patient(
        self,
        name: str,
        age: Any,
        gender: str,
        phone: str = "",
        email: str = "",
        address: str = "",
        summary: str = ""
    ) -> Dict[str, Any]:
        """
        Adds a new patient record to the Excel database or updates if already exists (deduplication).
        """
        wb = openpyxl.load_workbook(self.runtime_path)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=False))

        name_clean = str(name).strip()
        phone_clean = str(phone).strip()

        # Check if already exists to prevent duplicate rows
        existing_row_idx = None
        if len(rows) > 1:
            headers = [str(cell.value).strip() if cell.value is not None else "" for cell in rows[0]]
            name_col = headers.index("Name") if "Name" in headers else 2
            phone_col = headers.index("Phone_number") if "Phone_number" in headers else 0

            for r_idx in range(1, len(rows)):
                r = rows[r_idx]
                r_name = str(r[name_col].value or "").strip().lower()
                r_phone = str(r[phone_col].value or "").strip().replace("-", "").replace(" ", "").replace("+", "")
                p_clean = phone_clean.replace("-", "").replace(" ", "").replace("+", "")

                if (name_clean and name_clean.lower() == r_name) or (p_clean and len(p_clean) > 5 and p_clean == r_phone):
                    existing_row_idx = r_idx + 1
                    break

        if existing_row_idx:
            # Update existing row
            headers = [str(cell.value).strip() if cell.value is not None else "" for cell in rows[0]]
            for i, h in enumerate(headers):
                c = ws.cell(row=existing_row_idx, column=i + 1)
                if h == "Phone_number" and phone_clean:
                    c.value = phone_clean
                elif h == "Email" and email:
                    c.value = str(email).strip()
                elif h == "Name" and name_clean:
                    c.value = name_clean
                elif h == "Age" and age:
                    c.value = str(age).strip()
                elif h == "Gender" and gender:
                    c.value = str(gender).strip()
                elif h == "Address" and address:
                    c.value = str(address).strip()
                elif h == "Summary" and summary:
                    c.value = str(summary).strip()
        else:
            new_row = [
                str(phone).strip(),
                str(email).strip(),
                str(name).strip(),
                str(age).strip(),
                str(gender).strip(),
                str(address).strip(),
                str(summary).strip(),
            ]
            ws.append(new_row)

        wb.save(self.runtime_path)

        return {
            "Phone_number": str(phone),
            "Email": str(email),
            "Name": str(name),
            "Age": str(age),
            "Gender": str(gender),
            "Address": str(address),
            "Summary": str(summary)
        }

    def update_patient_summary(self, identifier: str, new_summary: str) -> bool:
        """
        Updates the Summary field for a patient matching identifier (name or phone).
        """
        wb = openpyxl.load_workbook(self.runtime_path)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=False))
        if not rows:
            return False

        headers = [str(cell.value).strip() if cell.value is not None else "" for cell in rows[0]]
        try:
            summary_col_idx = headers.index("Summary")
            name_col_idx = headers.index("Name")
            phone_col_idx = headers.index("Phone_number")
        except ValueError:
            return False

        target_row_idx = None
        ident_lower = identifier.lower().strip()
        clean_id = ident_lower.replace("-", "").replace(" ", "").replace("+", "")

        for r_idx in range(1, len(rows)):
            row = rows[r_idx]
            name_val = str(row[name_col_idx].value or "").lower()
            phone_val = str(row[phone_col_idx].value or "").lower().replace("-", "").replace(" ", "").replace("+", "")

            if ident_lower in name_val or name_val in ident_lower:
                target_row_idx = r_idx + 1
                break
            if clean_id and clean_id in phone_val:
                target_row_idx = r_idx + 1
                break

        if target_row_idx:
            ws.cell(row=target_row_idx, column=summary_col_idx + 1, value=str(new_summary).strip())
            wb.save(self.runtime_path)
            return True

        return False
